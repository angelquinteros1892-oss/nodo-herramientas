#!/usr/bin/env python3
"""
NODO HERRAMIENTAS - Importador de la lista maestra del proveedor.

Fuente maestra:
  ListaProductos.xlsx
Hoja:
  Productos - Servicios
Columnas:
  Categoria | Codigo | Nombre | $ Con IVA

La columna "Categoria" del proveedor NO se usa directamente como categoría
comercial de la tienda: suele representar marca/familia (LUSQTOFF, ORYX,
OMAHA, HOGAR Y BAZAR, etc.). La tienda deriva una categoria_tienda según
nombre/código para evitar que artículos válidos terminen en "Otros".
"""

import json
import os
import re
import shutil
import unicodedata
from datetime import datetime

import openpyxl

EXCEL_FILE = "ListaProductos.xlsx"
SHEET_NAME = "Productos - Servicios"
OUTPUT_FILE = "data/productos.json"
BACKUP_DIR = "data/backups"


def limpiar_texto(valor):
    if valor is None:
        return ""
    return str(valor).strip()


def normalizar(texto):
    texto = limpiar_texto(texto).lower()
    texto = unicodedata.normalize("NFD", texto)
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", texto)


REGLAS_CATEGORIA = [
    ("Soldadura", [
        "soldadora", "inverter", "electrodo", "mig", "tig", "mma",
        "mascara fotosensible", "careta", "pinza masa", "porta electrodo",
        "cincel sds",
    ]),
    ("Taladros y Atornilladores", [
        "taladro", "atornillador", "rotomartillo", "martillo demoledor",
        "percutor", "sds", "mecha", "broca",
    ]),
    ("Amoladoras y Corte", [
        "amoladora", "esmeril", "cortadora", "sierra circular", "sensitiva",
        "tronzadora", "disco de corte",
    ]),
    ("Carpintería y Madera", [
        "lijadora", "cepillo electrico", "fresadora", "caladora",
        "ingletadora", "sierra de banco",
    ]),
    ("Jardín y Exterior", [
        "motosierra", "desmalezadora", "bordeadora", "cortacerco",
        "cortadora de cesped", "sopladora", "podadora", "tijera de poda",
        "hidrolavadora", "fumigador",
    ]),
    ("Compresores y Neumática", [
        "compresor", "neumatic", "pistola para pintar", "pistola pintar",
        "inflador", "manguera aire", "acople rapido",
    ]),
    ("Generadores y Bombas", [
        "generador", "grupo electrogeno", "motobomba", "bomba de agua",
    ]),
    ("Automotor", [
        "crique", "arrancador", "cable puente", "cargador de bateria",
        "pulidora", "aspiradora auto", "llave impacto",
    ]),
    ("Herramientas Manuales", [
        "llave", "destornillador", "pinza", "alicate", "martillo",
        "prensa", "sargento", "tubo", "ratchet", "criquet", "cutter",
        "cinta metrica", "nivel", "serrucho",
    ]),
    ("Sets y Kits", [
        "set ", "kit ", "juego de ", "pack ", "caja de herramientas",
        "maletin", "valija",
    ]),
    ("Medición", [
        "tester", "multimetro", "laser", "medidor", "nivel laser",
        "termometro", "pinza amperometrica",
    ]),
    ("Iluminación", [
        "luz led", "reflector", "linterna", "luz de emergencia",
    ]),
    ("Hogar y Bazar", [
        "cafetera", "parlante", "calefactor", "ventilador", "pava",
        "aspiradora", "hidro", "cocina", "griferia",
    ]),
    ("Accesorios y Consumibles", [
        "disco", "lija", "electrodo", "boquilla", "accesorio",
        "cincel", "mecha", "broca", "bateria", "cargador",
    ]),
]


def clasificar_categoria(nombre, codigo="", categoria_fuente=""):
    texto = normalizar(f"{codigo} {nombre}")

    # Prioridad: categorías específicas por descripción.
    for categoria, palabras in REGLAS_CATEGORIA:
        if any(normalizar(p) in texto for p in palabras):
            return categoria

    # Fallbacks útiles según familia del proveedor.
    fuente = normalizar(categoria_fuente)
    if "griferia" in fuente:
        return "Hogar y Bazar"
    if "hogar" in fuente or "bazar" in fuente:
        return "Hogar y Bazar"

    return "Otros"


def exportar_productos():
    if not os.path.exists(EXCEL_FILE):
        raise FileNotFoundError(
            f"No se encontró {EXCEL_FILE}. Debe estar en la raíz del proyecto."
        )

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise RuntimeError(
            f"No existe la hoja '{SHEET_NAME}'. Hojas: {', '.join(wb.sheetnames)}"
        )

    ws = wb[SHEET_NAME]
    productos = []

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        categoria_fuente, codigo, nombre, precio_con_iva = row[:4]

        codigo = limpiar_texto(codigo)
        nombre = limpiar_texto(nombre)
        categoria_fuente = limpiar_texto(categoria_fuente)

        if not codigo and not nombre:
            continue

        try:
            precio = round(float(precio_con_iva or 0), 2)
        except (TypeError, ValueError):
            precio = 0

        productos.append({
            "id": i,
            "codigo": codigo or f"PROD-{i}",
            "nombre": nombre or codigo,
            "categoria": clasificar_categoria(nombre, codigo, categoria_fuente),
            "categoria_fuente": categoria_fuente,
            "precio_con_iva": precio,
            "precio_minorista": precio,
            "imagen": "",
            "activo": True,
        })

    os.makedirs(BACKUP_DIR, exist_ok=True)
    if os.path.exists(OUTPUT_FILE):
        fecha = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(
            OUTPUT_FILE,
            os.path.join(BACKUP_DIR, f"productos_{fecha}.json"),
        )

    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(productos, f, ensure_ascii=False, indent=2)

    conteo = {}
    for p in productos:
        conteo[p["categoria"]] = conteo.get(p["categoria"], 0) + 1

    print(f"Productos importados: {len(productos)}")
    for categoria, cantidad in sorted(conteo.items()):
        print(f"{categoria}: {cantidad}")

    otros = [p for p in productos if p["categoria"] == "Otros"]
    print(f"Pendientes de clasificar: {len(otros)}")

    return productos


if __name__ == "__main__":
    exportar_productos()
